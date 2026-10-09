"""文件解析子进程入口；接收受限字节，输出结构化正文或安全错误，不接触数据库和凭证。"""
import base64
import json
import re
import sys
from io import BytesIO
from zipfile import ZipFile, BadZipFile

MAX_CHARS = 1_000_000
MAX_PAGES = 300


class ParseError(Exception):
    """可返回给用户的格式或资源限制错误，不包含内部路径、SDK 栈或凭证。"""

    def __init__(self,code: str,message: str):
        """接收错误码与中文说明，保存可安全输出的解析错误，无返回值。"""
        super().__init__(message)
        self.code,self.message = code,message


def check_docx(data: bytes):
    """接收 DOCX 原件字节，检查压缩规模、文件结构及 XML 实体；合法时无返回值，非法抛安全错误。"""
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members)>5000 or sum(item.file_size for item in members)>100*1024*1024:
                raise ParseError('FILE_EXPANSION_LIMIT','DOCX 解压后内容过大，请缩小文件')
            if '[Content_Types].xml' not in archive.namelist() or 'word/document.xml' not in archive.namelist():
                raise ParseError('INVALID_DOCX','文件不是有效的 DOCX 文档')
            for item in members:
                if item.filename.endswith(('.xml','.rels')):
                    content = archive.read(item)
                    if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
                        raise ParseError('UNSAFE_XML','DOCX 包含不支持的 XML 实体')
    except BadZipFile:
        raise ParseError('INVALID_DOCX','DOCX 文件损坏或格式不正确')


def decode_text(data: bytes) -> str:
    """接收 TXT/Markdown 字节，以 UTF-8 或 GB18030 解码，返回正文；二进制和未知编码明确拒绝。"""
    if b'\0' in data:
        raise ParseError('BINARY_TEXT','文本文件含二进制内容，无法解析')
    for encoding in ('utf-8-sig','gb18030'):
        try:
            value = data.decode(encoding)
            if any(ord(char)<32 and char not in '\r\n\t' for char in value):
                raise ParseError('BINARY_TEXT','文本文件含不支持的控制字符')
            return value
        except UnicodeDecodeError:
            continue
    raise ParseError('TEXT_ENCODING','文本编码不支持，请保存为 UTF-8 后重试')


def docx_heading_level(paragraph) -> int | None:
    """接收真实 Word 段落，先读直接大纲级别再读样式继承，返回1～9级标题；正文级别9不当作标题。"""
    values=paragraph._p.xpath('./w:pPr/w:outlineLvl/@w:val')
    if values:
        return int(values[0])+1 if values[0].isdigit() and 0<=int(values[0])<=8 else None
    style=paragraph.style
    visited=set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        values=style.element.xpath('./w:pPr/w:outlineLvl/@w:val')
        if values:
            return int(values[0])+1 if values[0].isdigit() and 0<=int(values[0])<=8 else None
        match=re.fullmatch(r'(?:Heading|标题)\s*([1-9])',style.name or '',re.I)
        if match:
            return int(match[1])
        if (style.name or '').lower() in ('title','标题'):
            return 1
        style=style.base_style
    return None


def parse(data: bytes,kind: str) -> dict:
    """接收真实原件与服务端格式，返回完整受限正文、段落/表格或 PDF 页；不生成任何正式知识切片。"""
    if kind in ('txt','md'):
        text = decode_text(data)
        result = {'format':kind,'page_count':1,'blocks':[{'type':'text','text':text}],'text':text}
    elif kind=='docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        check_docx(data)
        blocks,texts = [],[]
        for item in Document(BytesIO(data)).iter_inner_content():
            if isinstance(item,Paragraph):
                if item.text:
                    blocks.append({'type':'paragraph','text':item.text,'style':item.style.name if item.style else '',
                                   'heading_level':docx_heading_level(item)})
                    texts.append(item.text)
            else:
                rows = [[cell.text for cell in row.cells] for row in item.rows]
                blocks.append({'type':'table','rows':rows})
                texts.extend('\t'.join(row) for row in rows)
            if sum(len(value) for value in texts)>MAX_CHARS:
                raise ParseError('TEXT_LIMIT','文档正文过长，请分成较小文件')
        result = {'format':kind,'page_count':1,'blocks':blocks,'text':'\n'.join(texts)}
    elif kind=='pdf':
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(data) as document:
            if len(document)>MAX_PAGES:
                raise ParseError('PAGE_LIMIT','PDF 超过 300 页，请分批提交')
            pages,texts = [],[]
            for index in range(len(document)):
                page = document[index]
                try:
                    textpage = page.get_textpage()
                    try:
                        value = textpage.get_text_range()
                    finally:
                        textpage.close()
                finally:
                    page.close()
                pages.append({'page':index+1,'text':value})
                texts.append(value)
                if sum(len(text) for text in texts)>MAX_CHARS:
                    raise ParseError('TEXT_LIMIT','PDF 正文过长，请分成较小文件')
            result = {'format':kind,'page_count':len(document),'pages':pages,'text':'\n'.join(texts)}
    else:
        raise ParseError('UNSUPPORTED_FORMAT','当前仅支持 PDF、DOCX、TXT 和 Markdown')
    if len(result['text'])>MAX_CHARS:
        raise ParseError('TEXT_LIMIT','正文超过处理限制，请缩小文件')
    if not result['text'].strip():
        raise ParseError('EMPTY_TEXT','未提取到正文；扫描 PDF 和纯图片文档暂不支持 OCR')
    return result


def render(data: bytes,page_number: int) -> bytes:
    """接收 PDF 字节和从 1 开始页码，返回限定像素的 PNG 页面；不执行 PDF 脚本或发送外部请求。"""
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(data) as document:
        if not 1<=page_number<=len(document) or len(document)>MAX_PAGES:
            raise ParseError('INVALID_PAGE','PDF 页码不正确或页数超过限制')
        page = document[page_number-1]
        try:
            width,height = page.get_size()
            if width<=0 or height<=0:
                raise ParseError('INVALID_PAGE','PDF 页面尺寸不正确')
            scale = min(1.5,(2_000_000/(width*height))**0.5)
            bitmap = page.render(scale=scale)
            try:
                image = bitmap.to_pil()
                output = BytesIO()
                image.save(output,format='PNG')
                image.close()
                return output.getvalue()
            finally:
                bitmap.close()
        finally:
            page.close()


def main():
    """无参数读取 stdin 的受限 JSON，执行解析/渲染并仅输出结果；Linux 子进程限制内存与 CPU。"""
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS,(768*1024*1024,768*1024*1024))
        resource.setrlimit(resource.RLIMIT_CPU,(30,30))
    except ImportError:
        pass
    try:
        raw = sys.stdin.buffer.read(80*1024*1024+1)
        if len(raw)>80*1024*1024:
            raise ParseError('FILE_LIMIT','原件过大')
        request = json.loads(raw)
        data = base64.b64decode(request['data'],validate=True)
        if request.get('operation')=='render':
            value = {'png':base64.b64encode(render(data,request['page'])).decode()}
        else:
            value = parse(data,request['format'])
        print(json.dumps({'ok':True,'result':value},ensure_ascii=False))
    except Exception as error:
        code = error.code if isinstance(error,ParseError) else 'INVALID_FILE'
        message = error.message if isinstance(error,ParseError) else '文件损坏、受密码保护或格式不正确，无法解析'
        print(json.dumps({'ok':False,'error':{'code':code,'message':message}},ensure_ascii=False))


if __name__=='__main__':
    main()
