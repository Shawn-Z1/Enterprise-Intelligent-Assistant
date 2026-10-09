import json
import unittest
from io import BytesIO
from pathlib import Path

import test_m06 as hybrid
from app.chat_worker import queue_existing_approved
from app.ids import new_id
from app.knowledge_chunking import looks_like_heading
from app.rag_contract import INDEX_VERSION


class ChunkingIntegration(unittest.IsolatedAsyncioTestCase):
    """真实 DOCX/Markdown/PDF 原件、PG 元数据、CPU 全窗口模型及 Milvus 索引的切片验收。"""
    login_as=hybrid.RagTests.login_as
    call=hybrid.RagTests.call
    headers=hybrid.RagTests.headers
    upload=hybrid.RagTests.upload
    submitted=hybrid.RagTests.submitted
    review=hybrid.RagTests.review
    drain=hybrid.RagTests.drain
    context=hybrid.RagTests.context
    indexed=hybrid.RagTests.indexed
    search=hybrid.RagTests.search
    asyncSetUp=hybrid.RagTests.asyncSetUp
    asyncTearDown=hybrid.RagTests.asyncTearDown

    async def test_real_docx_hierarchy_metadata_and_search(self):
        """真实规范 DOCX 按长章/节下钻，正文与标题一起入库，PG 和 Milvus 均保留祖先路径且实际查询命中。"""
        session=self.login_as('admin')
        source=Path('/tests/fixtures/chunking/hierarchy.docx').read_bytes()
        entry=self.submitted(session,[('hierarchy.docx',source)])[0]
        await self.drain()
        rows=[dict(row) for row in await self.admin.fetch('SELECT * FROM app_knowledge.chunks WHERE status=\'active\' ORDER BY position')]
        self.assertEqual(len(rows),4)
        self.assertTrue(all(len(row['content'])<=600 and row['index_version']==INDEX_VERSION and not looks_like_heading(row['body_content']) for row in rows))
        software=[row for row in rows if any(value['title']=='8.2 软件要求' for value in row['heading_path'])]
        self.assertEqual(len(software),2)
        self.assertTrue(all(len(row['heading_path'])==3 for row in software))
        self.assertTrue(all('第八章 技术要求' in row['content'] and '8.2 软件要求' in row['content'] for row in software))
        result=await self.search(self.login_as('zhangsan'),'ZX-CHUNK-A17 软件环境的数据库版本要求是什么？')
        self.assertTrue(result['citations'])
        self.assertTrue(any('PostgreSQL 17' in row['content'] for row in result['citations']))
        records=self.index.client.query(self.index.collection,filter='',output_fields=['chunk_id','heading_path','page_numbers'],limit=10)
        self.assertEqual(len(records),4)
        self.assertTrue(any(len(row['heading_path'])==3 for row in records))
        selected=result['citations'][0]
        response=self.call(self.login_as('zhangsan'),'GET','/knowledge/sources/'+selected['chunk_id'])
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['heading_path'],selected['heading_path'])

    async def test_real_markdown_top_sections_and_pdf_cross_page(self):
        """真实 Markdown 保持两章边界；短 PDF 两页正文合并成一片，来源接口准确显示页集合。"""
        admin=self.login_as('admin')
        markdown='# 第一章\n第一章制度正文。\n## 小节\n小节正文。\n# 第二章\n第二章业务正文。'
        entry=self.submitted(admin,[('chapters.md',markdown.encode())])[0]
        await self.drain()
        rows=await self.admin.fetch('SELECT k.* FROM app_knowledge.chunks k JOIN app_knowledge.documents d USING(document_id) WHERE d.application_id=$1',entry['application_id'])
        self.assertEqual(len(rows),2)
        pdf=self.submitted(admin,[('travel.pdf',Path('/tests/fixtures/m06/travel.pdf').read_bytes())])[0]
        await self.drain()
        rows=await self.admin.fetch('SELECT k.* FROM app_knowledge.chunks k JOIN app_knowledge.documents d USING(document_id) WHERE d.application_id=$1',pdf['application_id'])
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['page_numbers'],[1,2])
        response=self.call(admin,'GET','/knowledge/sources/'+rows[0]['chunk_id'])
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['page_numbers'],[1,2])

    async def test_word_outline_without_named_heading_style(self):
        """真实 Word 原件只有直接大纲级别时，解析和最终切片仍保存三级祖先，复现用户文件格式。"""
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        document=Document()
        for level,title in ((1,'系统介绍'),(2,'八、部署环境要求'),(3,'8.1 硬件要求')):
            paragraph=document.add_paragraph(title)
            element=OxmlElement('w:outlineLvl');element.set(qn('w:val'),str(level-1));paragraph._p.get_or_add_pPr().append(element)
        document.add_paragraph('硬件应满足必要要求。'*40)
        paragraph=document.add_paragraph('8.2 软件要求')
        element=OxmlElement('w:outlineLvl');element.set(qn('w:val'),'2');paragraph._p.get_or_add_pPr().append(element)
        document.add_paragraph('软件应使用 PostgreSQL 17。'*30)
        stream=BytesIO();document.save(stream)
        admin=self.login_as('admin')
        self.submitted(admin,[('outline.docx',stream.getvalue())])
        await self.drain()
        rows=await self.admin.fetch("SELECT heading_path,content FROM app_knowledge.chunks WHERE status='active'")
        software=[row for row in rows if row['heading_path'][-1]['title']=='8.2 软件要求']
        self.assertTrue(software)
        self.assertTrue(all([item['title'] for item in row['heading_path']]==['系统介绍','八、部署环境要求','8.2 软件要求'] for row in software))
        self.assertTrue(all(len(row['content'])<=600 for row in rows))

    async def test_policy_upgrade_reindexes_and_deleted_documents_stay_deleted(self):
        """旧版批准文档自动排重建，保留旧切片/审批，建立独立新索引；已失效文档绝不重新激活。"""
        entry=await self.indexed(member=True)
        document=entry['document_id']
        old_collection=self.name+'_legacy'
        self.index.client.create_collection(old_collection,dimension=512,metric_type='COSINE')
        self.index.client.insert(old_collection,[{'id':1,'vector':[1.0]+[0.0]*511,'document_id':document}])
        await self.admin.execute('UPDATE app_knowledge.documents SET embedding_space=\'legacy-space\',index_version=1,collection_name=$1 WHERE document_id=$2',old_collection,document)
        await self.admin.execute('UPDATE app_knowledge.chunks SET embedding_space=\'legacy-space\',index_version=1 WHERE document_id=$1',document)
        # 清除本测试已建立的当前派生实体，模拟真正只有旧版本向量的待升级状态。
        await self.index.delete_document(document)
        await queue_existing_approved(self.pool,self.processor)
        self.assertEqual(await self.admin.fetchval('SELECT count(*) FROM app_knowledge.jobs WHERE job_type=\'reindex\''),1)
        await self.drain()
        self.assertEqual(await self.admin.fetchval('SELECT status FROM app_knowledge.documents'),'retrievable')
        self.assertEqual(await self.admin.fetchval('SELECT index_version FROM app_knowledge.documents'),INDEX_VERSION)
        self.assertEqual(await self.admin.fetchval("SELECT count(*) FROM app_knowledge.chunks WHERE index_version=1 AND status='superseded'"),1)
        self.assertEqual(await self.admin.fetchval("SELECT count(*) FROM app_knowledge.chunks WHERE index_version=$1 AND status='active'",INDEX_VERSION),1)
        self.assertEqual(self.index.client.query(old_collection,filter='document_id == '+json.dumps(document),output_fields=['id'],consistency_level='Strong'),[])
        self.assertEqual(await self.admin.fetchval('SELECT count(*) FROM app_knowledge.reviews'),1)
        self.assertTrue((await self.search(self.login_as('zhangsan')))['citations'])
        admin=self.login_as('admin')
        self.assertEqual(self.call(admin,'POST',f"/knowledge/applications/{entry['application_id']}/delete").status_code,200)
        await self.drain()
        await queue_existing_approved(self.pool,self.processor)
        self.assertEqual(await self.admin.fetchval("SELECT count(*) FROM app_knowledge.jobs WHERE status='pending'"),0)
        self.assertEqual(await self.admin.fetchval('SELECT status FROM app_knowledge.documents'),'invalidated')

    async def test_long_chinese_tail_is_embedded_and_reranked_without_truncation(self):
        """超过512token的真实中文输入处理全部窗口；尾部内容改变实际向量，输出仍是512维同空间单位向量。"""
        await self.index.ensure()
        filler=''.join(chr(0x4e00+index) for index in range(550))
        first=await self.index.embed([filler+'软件部署需要 PostgreSQL 17。'])
        second=await self.index.embed([filler+'财务报销需要电子发票。'])
        self.assertGreater(first['window_counts'][0],1)
        self.assertFalse(first['truncated'])
        self.assertEqual(len(first['values'][0]),512)
        self.assertNotEqual(first['values'][0],second['values'][0])
        scores=await self.index.model('/rerank',{'texts':['汉'*580+'软件部署需要 PostgreSQL 17。'],'query':'软件部署需要什么数据库版本？'})
        self.assertGreater(scores['window_counts'][0],1)
        self.assertFalse(scores['truncated'])


if __name__=='__main__':unittest.main()
