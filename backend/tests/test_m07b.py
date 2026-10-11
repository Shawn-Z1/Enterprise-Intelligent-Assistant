"""M07B目标验收：独立PG、虚构文件、真实Docker/MinIO与实际主审核模型；不做无目标压力测试。"""
import asyncio
import base64
import hashlib
import io
import json
import unittest
from pathlib import Path
from types import SimpleNamespace

import httpx
from docx import Document
import test_m04 as base
from test_m07a import CSV, workbook
from app.config import Settings
from app.chat_service import ChatProcessor, RunWriter
from app.knowledge_storage import KnowledgeStorage
from app.model_gateway import ModelGateway
from app.ids import new_id
from app.rag_service import live_context
from app.sandbox_service import SandboxArguments, SandboxPending, SandboxService
from app.sandbox_queue import recover_queue


class Runtime(unittest.IsolatedAsyncioTestCase):
    """直接验证可信执行器与所有声明格式，固定测试代码仍只在独立任务容器执行。"""
    async def asyncSetUp(self):
        """读取控制配置并创建受限诊断客户端，不读取业务表；返回准备好的测试上下文。"""
        self.settings=Settings();self.client=httpx.AsyncClient(timeout=85,trust_env=False)
        self.headers={'Authorization':'Bearer '+self.settings.sandbox_control_token.get_secret_value()};self.evidence=[]

    async def asyncTearDown(self):
        """保存本方法真实回执并关闭HTTP，不删除原件或其他人的节点状态。"""
        folder=Path('/verification/M07B');folder.mkdir(parents=True,exist_ok=True)
        (folder/(self._testMethodName+'.json')).write_text(json.dumps(self.evidence,ensure_ascii=False,indent=2),encoding='utf-8')
        await self.client.aclose()

    def file(self, raw, kind, name=None):
        """接收虚构文件字节/类型，生成真实哈希与仅供测试的附件清单，不含宿主路径。"""
        return {'attachment_id':new_id(),'name':name or 'fixture.'+kind,'format':kind,'sha256':hashlib.sha256(raw).hexdigest(),'data':base64.b64encode(raw).decode()}

    async def run_code(self, code='', files=None, operation='execute', language='python', read=None):
        """接收固定测试任务，调用真实私有Runner执行并返回回执；不冒充模型生成的代码。"""
        packet={'user_id':new_id(),'operation':operation,'language':language,'code':code,'files':files or [],'read':read}
        response=await self.client.post(self.settings.sandbox_runner_url+'/jobs/'+new_id(),headers=self.headers,json=packet)
        self.assertEqual(response.status_code,200,response.text);value=response.json();self.evidence.append({'operation':operation,'receipt':value})
        return value

    async def test_formats_paging_conversion_and_xlsx(self):
        """固定输入覆盖六格式、分页尾部、真实DOCX/PDF转换与既有DuckDB/XLSX路径。"""
        doc=Document();doc.add_heading('虚构会议',0);doc.add_paragraph('李明负责周五交付，预算120。')
        stream=io.BytesIO();doc.save(stream)
        files=[self.file(b'# Meeting\nOwner: Alice, deadline Friday.','md'),self.file('中文文本'.encode(),'txt'),
               self.file(CSV,'csv'),self.file(workbook(),'xlsx'),self.file(stream.getvalue(),'docx')]
        prepared=await self.run_code(files=files,operation='prepare');self.assertTrue(prepared['ok'],prepared)
        self.assertTrue(all(x['status']=='ready' for x in prepared['result']['files']),prepared)
        converted=await self.run_code("convert_document(next(iter(files.values())), 'pdf', 'meeting')",[files[0]])
        self.assertTrue(converted['ok'],converted);self.assertEqual({x['name'] for x in converted['artifacts']},{'meeting.docx','meeting.pdf'})
        pdf=next(x for x in converted['artifacts'] if x['name'].endswith('.pdf'))
        parsed=await self.run_code(files=[self.file(base64.b64decode(pdf['data']),'pdf')],operation='prepare')
        self.assertTrue(parsed['ok']);self.assertEqual(parsed['result']['files'][0]['status'],'ready')
        self.assertIn('Meeting',parsed['result']['files'][0]['sample'])
        long=self.file(('开始\n'+'记录内容\n'*700+'尾部负责人：赵明').encode(),'txt')
        overview=await self.run_code(files=[long],operation='prepare');self.assertFalse(overview['result']['files'][0]['sample_complete'])
        page=await self.run_code(files=[long],operation='read',read={'attachment_id':long['attachment_id'],'offset':3300,'limit':12000})
        self.assertTrue(page['ok']);self.assertIn('赵明',page['result']['text']);self.assertFalse(page['result']['has_more'])
        sql=await self.run_code('SELECT count(*) AS n FROM f1_s1',[files[3]],language='sql')
        self.assertTrue(sql['ok'],sql);self.assertGreater(sql['result']['rows'][0]['n'],0)
        wide=self.file((','.join('c'+str(i) for i in range(25))+'\n'+','.join(str(i) for i in range(25))+'\n').encode(),'csv')
        first=await self.run_code(files=[wide],operation='read',read={'attachment_id':wide['attachment_id']})
        self.assertTrue(first['result']['has_more_columns']);self.assertFalse(first['result']['coverage']['complete'])
        last=await self.run_code(files=[wide],operation='read',read={'attachment_id':wide['attachment_id'],'column_offset':20})
        self.assertEqual(last['result']['rows'][0]['c24'],24);self.assertFalse(last['result']['has_more_columns'])

    async def test_invalid_files_and_isolation(self):
        """代表假文档/无文字PDF、跨实例文件/网络边界和非法产物；复用M07A其余资源验收。"""
        value=await self.run_code(files=[self.file(b'not a document','docx')],operation='prepare')
        self.assertEqual(value['result']['files'][0]['status'],'failed')
        blank=await self.run_code("from pypdf import PdfWriter\nw=PdfWriter();w.add_blank_page(width=100,height=100);w.write(str(OUTPUT_DIR/'blank.pdf'))")
        self.assertTrue(blank['ok'],blank)
        value=await self.run_code(files=[self.file(base64.b64decode(blank['artifacts'][0]['data']),'pdf')],operation='prepare')
        self.assertIn('OCR',value['result']['files'][0]['error'])
        value=await self.run_code("from pathlib import Path\nPath('/tmp/marker').write_text('private')\nemit_text(str(Path('/var/run/docker.sock').exists()))")
        self.assertTrue(value['ok']);self.assertEqual(value['result']['text'],'False')
        value=await self.run_code("from pathlib import Path\nassert not Path('/tmp/marker').exists()\n(OUTPUT_DIR/'bad.png').symlink_to('/etc/passwd')")
        self.assertFalse(value['ok']);self.assertFalse(value['artifacts'])
        value=await self.run_code("(OUTPUT_DIR/'fake.pdf').write_bytes(b'%PDF-fake %%EOF')")
        self.assertFalse(value['ok']);self.assertFalse(value['artifacts'])


class FileFlow(unittest.IsolatedAsyncioTestCase):
    """使用既有隔离库准备助手但不继承其全部旧测试；真实成功链路采用实际主模型和审核模型。"""
    login=base.ChatTests.login

    async def asyncSetUp(self):
        """建立本方法独立PG及真实MinIO/Runner，测试模型配置不读取主库企业资料。"""
        await base.ChatTests.asyncSetUp(self);self.actual=Settings();self.evidence=[]
        names=('sandbox_runner_url','sandbox_control_token','app_minio_endpoint','app_minio_access_key','app_minio_secret_key','app_minio_bucket',
            'primary_base_url','primary_api_key','primary_model','fallback_base_url','fallback_api_key','fallback_model','review_base_url','review_api_key','review_model')
        self.settings=self.settings.model_copy(update={name:getattr(self.actual,name) for name in names}|{'sandbox_enabled':True,'primary_context_tokens':65536,'fallback_context_tokens':65536,'review_context_tokens':65536})
        self.api.app.state.settings=self.settings;self.storage=KnowledgeStorage(self.settings);self.api.app.state.knowledge_storage=self.storage
        self.node=httpx.AsyncClient(timeout=85,trust_env=False);self.control_headers={'Authorization':'Bearer '+self.actual.sandbox_control_token.get_secret_value()}

    async def asyncTearDown(self):
        """只删除本测试数据库/对象，保存无密钥证据；节点控制诊断不修改其他任务。"""
        folder=Path('/verification/M07B');folder.mkdir(parents=True,exist_ok=True)
        (folder/(self._testMethodName+'.json')).write_text(json.dumps(self.evidence,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
        rows=await self.admin.fetch('SELECT object_key FROM app_runtime.analysis_attachments UNION ALL SELECT object_key FROM app_runtime.sandbox_artifacts')
        for row in rows:await self.storage.remove_new(row['object_key'])
        await self.node.aclose();await base.ChatTests.asyncTearDown(self)

    def upload(self, headers, files):
        """接收虚构文件列表，经过真实会话/multipart/CSRF路由上传，返回当前会话和附件编号。"""
        response=self.api.post('/api/conversations',headers=headers,json={});self.assertEqual(response.status_code,201,response.text)
        conversation=response.json()['conversation_id']
        response=self.api.post('/api/conversations/'+conversation+'/attachments',headers=headers,files=[('files',(name,data,'application/octet-stream')) for name,data in files])
        self.assertEqual(response.status_code,201,response.text);return conversation,[x['attachment_id'] for x in response.json()]

    def processor(self):
        """创建实际模型网关及生产Agent；此文件场景不走企业检索，不提供假RAG成功响应。"""
        gateway=ModelGateway(self.settings);self.gateways.append(gateway)
        return ChatProcessor(self.pool,self.settings,gateway,self.audit,rag=SimpleNamespace())

    async def claim_job(self, capacity=2):
        """测试调度驱动调用生产领取接口；返回可信packet，真实用户路径由Runner后台自动驱动。"""
        health=(await self.node.get(self.actual.sandbox_runner_url+'/health',headers=self.control_headers)).json()
        response=self.api.post('/internal/sandbox/claim',headers=self.control_headers,json={'runner_id':'sandbox-runner-1','capacity':capacity,'image_id':health['image_id']})
        self.assertEqual(response.status_code,200,response.text);return response.json()['job']

    async def complete_job(self, job):
        """接收生产队列领取包，实际跑独立Docker，再调用生产持久回执接口，不用固定假结果。"""
        response=await self.node.post(self.actual.sandbox_runner_url+'/jobs/'+job['job_id'],headers=self.control_headers,json=job)
        self.assertEqual(response.status_code,200,response.text);receipt=response.json()
        response=self.api.post('/internal/sandbox/jobs/'+job['job_id']+'/complete',headers=self.control_headers,json={'runner_id':'sandbox-runner-1','lease_version':job['lease_version'],'receipt':receipt})
        self.assertEqual(response.status_code,200,response.text);self.assertTrue(response.json()['saved'])
        self.evidence.append({'job_id':job['job_id'],'operation':job['operation'],'receipt_ok':receipt['ok']})

    async def flow(self, files, question):
        """运行实际模型/强制预处理/持久等待/容器/审核，返回真实答案与调用和产物记录。"""
        headers=self.login();conversation,ids=self.upload(headers,files);processor=self.processor()
        response=self.api.post('/api/conversations/'+conversation+'/messages',headers=headers,json={'content':question,'attachment_ids':ids,'request_key':new_id()})
        self.assertEqual(response.status_code,202,response.text);ident=response.json()['run_id']
        first=await processor.claim();await processor.process(first)
        state=await self.admin.fetchrow('SELECT * FROM app_runtime.runs WHERE run_id=$1',ident)
        self.assertEqual(state['status'],'waiting_sandbox',state['state_json']);self.assertIsNone(state['lease_owner'])
        self.assertEqual(await self.admin.fetchval('SELECT count(*) FROM app_runtime.model_attempts WHERE trace_id=$1',state['trace_id']),0)
        for _ in range(35):
            job=await self.claim_job()
            if job:await self.complete_job(job)
            run=await processor.claim()
            if run:await processor.process(run)
            state=await self.admin.fetchrow('SELECT * FROM app_runtime.runs WHERE run_id=$1',ident)
            if state['status'] in ('completed','failed','cancelled','interrupted'):break
            await asyncio.sleep(.05)
        self.assertEqual(state['status'],'completed',state['state_json'])
        self.assertEqual(await self.admin.fetchval('SELECT count(*) FROM app_runtime.query_approvals WHERE run_id=$1',ident),0)
        reply=await self.admin.fetchrow("SELECT content,protocol_json FROM app_runtime.messages WHERE run_id=$1 AND role='assistant'",ident)
        records=await self.admin.fetch('SELECT tool_name,input_json,output_json,status FROM app_runtime.tool_calls WHERE run_id=$1 ORDER BY created_at',ident)
        artifacts=await self.admin.fetch('SELECT a.* FROM app_runtime.sandbox_artifacts a JOIN app_runtime.sandbox_results r USING(result_id) WHERE r.job_id IN(SELECT job_id FROM app_runtime.sandbox_jobs WHERE run_id=$1)',ident)
        self.evidence.append({'run_id':ident,'question':question,'answer':dict(reply),'tools':[dict(x) for x in records],'artifacts':[dict(x) for x in artifacts]})
        self.assertEqual(reply['protocol_json']['message']['answer_review']['status'],'approved',reply['content'])
        return reply,records,artifacts,headers

    async def test_real_text_and_full_document(self):
        """真实主审核模型完成短纪要及长DOCX尾部读取，预处理必须先于首次模型请求。"""
        reply,records,_,_=await self.flow([('meeting.md','# 虚构例会\n周五由林明交付测试报告。预算120元。上线日期尚未确定。'.encode())],
            '仅依据这份虚构会议文字整理纪要，列出待办、负责人和期限，未确定的内容标未确定。不需要生成文件。')
        self.assertIn('林明',reply['content']);self.assertFalse(any(x['tool_name']=='run_sandbox_code' for x in records))
        doc=Document();doc.add_heading('虚构项目会议',0)
        for i in range(220):doc.add_paragraph(f'第{i+1}项进度记录：功能正在验证，风险尚待讨论。')
        doc.add_paragraph('最后决议：负责人程远，交付日期2026年10月20日，预算730元。')
        stream=io.BytesIO();doc.save(stream)
        reply,records,_,_=await self.flow([('long_meeting.docx',stream.getvalue())],
            '请完整读取这份虚构Word会议资料，尤其核对末尾决议，给出最后确定的负责人、交付日期和预算。不要根据概况样例猜测。')
        self.assertTrue(any(x['tool_name']=='read_attachment' for x in records));self.assertIn('程远',reply['content']);self.assertIn('730',reply['content'])

    async def test_real_multi_csv_heatmap_and_conversion(self):
        """实际模型生成多CSV联查与热力图，再实际MD转DOCX/PDF；产物下载SHA和另一用户边界真实验证。"""
        orders=b'person,month,amount\na,9,120\na,10,80\nb,9,200\nb,10,100\n'
        people='person,team\na,平台组\nb,数据组\n'.encode()
        reply,records,artifacts,headers=await self.flow([('orders.csv',orders),('people.csv',people)],
            '仅使用两个虚构CSV，按person联查，统计各team各month金额并生成热力图PNG与汇总CSV。请通过Python实际计算，说明合计，不需要联网或知识库。')
        self.assertTrue(any(x['tool_name']=='run_sandbox_code' and x['status']=='completed' for x in records))
        self.assertTrue(any(x['filename'].endswith('.png') for x in artifacts));self.assertIn('500',reply['content'])
        for artifact in artifacts:
            response=self.api.get('/api/sandbox/artifacts/'+artifact['artifact_id'],headers=headers);self.assertEqual(response.status_code,200)
            self.assertEqual(hashlib.sha256(response.content).hexdigest(),artifact['sha256'])
        other=self.login('lisi');self.assertEqual(self.api.get('/api/sandbox/artifacts/'+artifacts[0]['artifact_id'],headers=other).status_code,404)
        reply,records,artifacts,_=await self.flow([('report.md','# 虚构报告\n\n预算730元，负责人程远。\n\n|事项|状态|\n|---|---|\n|测试|完成|'.encode())],
            '将这份Markdown原件转换为可下载的Word DOCX和PDF，保持标题和表格。请实际运行转换并提供两个真实文件，不要仅回复文字。')
        self.assertTrue({'docx','pdf'}<={x['filename'].rsplit('.',1)[-1] for x in artifacts})

    async def test_queue_wait_and_cancel(self):
        """单槽实际任务领取/等待/取消与自动续取，不做大规模压力。"""
        headers=self.login();conversation,ids=self.upload(headers,[('data.csv',CSV)])
        response=self.api.post('/api/conversations/'+conversation+'/messages',headers=headers,json={'content':'虚构队列验证','attachment_ids':ids,'request_key':new_id()})
        ident=response.json()['run_id'];processor=self.processor();run=await processor.claim();writer=RunWriter(self.pool,self.audit,run)
        context=await live_context(self.pool,run['user_id'],run['state_json']['auth_session_id'])
        async def submit(label, task_run=run, task_context=context, task_writer=writer, attachments=ids):
            """接收新调用标签，注册真实工具记录并提交固定测试代码，仅用于调度边界而非模型正路径。"""
            call_id=new_id();await self.admin.execute("INSERT INTO app_runtime.tool_calls(call_id,run_id,tool_name,input_json,status,idempotency_key) VALUES($1,$2,'run_sandbox_code','{}','running',$3)",call_id,task_run['run_id'],task_run['run_id']+label)
            with self.assertRaises(SandboxPending) as waiting:await processor.sandbox.execute(SandboxArguments(attachment_ids=attachments,code='import time\ntime.sleep(1)\nemit_result([])',purpose=label),task_context,task_run,call_id,task_writer)
            return waiting.exception.job_id
        first=await submit('first')
        other=self.login('lisi');other_conversation,other_ids=self.upload(other,[('other.csv',CSV)])
        self.api.post('/api/conversations/'+other_conversation+'/messages',headers=other,json={'content':'另一用户队列验证','attachment_ids':other_ids,'request_key':new_id()})
        other_run=await processor.claim();other_context=await live_context(self.pool,other_run['user_id'],other_run['state_json']['auth_session_id'])
        second=await submit('second',other_run,other_context,RunWriter(self.pool,self.audit,other_run),other_ids)
        job=await self.claim_job(capacity=1);self.assertEqual(job['job_id'],first)
        self.assertIsNone(await self.claim_job(capacity=1))
        await self.complete_job(job);job=await self.claim_job(capacity=1);self.assertEqual(job['job_id'],second);await self.complete_job(job)
        third=await submit('third');self.api.post('/api/runs/'+ident+'/cancel',headers=headers)
        await recover_queue(self.pool,self.audit)
        self.assertEqual(await self.admin.fetchval('SELECT status FROM app_runtime.sandbox_jobs WHERE job_id=$1',third),'cancelled')
        self.assertIsNone(await self.claim_job())

    async def test_changed_input_unknown_and_attempt_budget(self):
        """排队中原件变化不启动代码；失效租约标未知不重跑；自动修正仍受三次尝试约束。"""
        headers=self.login();conversation,ids=self.upload(headers,[('fixture.csv',CSV)])
        queued=self.api.post('/api/conversations/'+conversation+'/messages',headers=headers,json={'content':'虚构故障验证','attachment_ids':ids,'request_key':new_id()}).json()
        processor=self.processor();run=await processor.claim();writer=RunWriter(self.pool,self.audit,run)
        context=await live_context(self.pool,run['user_id'],run['state_json']['auth_session_id'])
        args=SandboxArguments(attachment_ids=ids,code="emit_result([{'value':42}])",purpose='固定故障测试')
        async def submit():
            """注册新的固定测试调用并提交，返回call/job；测试不使用假模型或假执行结果。"""
            call=new_id();await self.admin.execute("INSERT INTO app_runtime.tool_calls(call_id,run_id,tool_name,input_json,status,idempotency_key) VALUES($1,$2,'run_sandbox_code','{}','running',$3)",call,run['run_id'],new_id())
            with self.assertRaises(SandboxPending) as waiting:await processor.sandbox.execute(args,context,run,call,writer)
            return call,waiting.exception.job_id
        call,ident=await submit();file=await self.admin.fetchrow('SELECT * FROM app_runtime.analysis_attachments WHERE attachment_id=$1',ids[0])
        await self.storage.put(file['object_key'],b'changed input')
        self.assertIsNone(await self.claim_job())
        state=await self.admin.fetchrow('SELECT status,result_json,started_at FROM app_runtime.sandbox_jobs WHERE job_id=$1',ident)
        self.assertEqual(state['status'],'failed');self.assertEqual(state['result_json']['error']['code'],'ATTACHMENT_EVIDENCE_CHANGED')
        await self.storage.put(file['object_key'],CSV)
        call,ident=await submit();job=await self.claim_job();self.assertEqual(job['job_id'],ident)
        await self.admin.execute("UPDATE app_runtime.sandbox_jobs SET lease_until=now()-interval '1 second' WHERE job_id=$1",ident)
        await recover_queue(self.pool,self.audit)
        self.assertEqual(await self.admin.fetchval('SELECT status FROM app_runtime.sandbox_jobs WHERE job_id=$1',ident),'unknown')
        repeated=await processor.sandbox.execute(args,context,run,call,writer)
        self.assertEqual(repeated['error']['code'],'SANDBOX_RESULT_UNKNOWN');self.assertIsNone(await self.claim_job())
        _,third=await submit();await self.complete_job(await self.claim_job())
        call=new_id();await self.admin.execute("INSERT INTO app_runtime.tool_calls(call_id,run_id,tool_name,input_json,status,idempotency_key) VALUES($1,$2,'run_sandbox_code','{}','running',$3)",call,run['run_id'],new_id())
        with self.assertRaises(Exception) as limit:await processor.sandbox.execute(args,context,run,call,writer)
        self.assertEqual(limit.exception.detail['code'],'SANDBOX_ATTEMPT_LIMIT')
        self.assertEqual(self.api.post('/internal/sandbox/claim',json={'runner_id':'forged'}).status_code,401)
        unsupported=self.api.post('/api/conversations/'+conversation+'/attachments',headers=headers,files=[('files',('image.png',b'not an allowed attachment'))])
        self.assertEqual(unsupported.status_code,422)
        unavailable=SandboxService(self.pool,self.settings.model_copy(update={'sandbox_runner_url':self.actual.sandbox_runner_url+'/unavailable'}),self.audit,None)
        with self.assertRaises(Exception) as missing:await unavailable.health()
        self.assertEqual(missing.exception.detail['code'],'SANDBOX_UNAVAILABLE')
        self.evidence.append({'input_changed_no_execution':True,'unknown_not_reexecuted':True,'attempt_budget':3,'fourth_rejected':True})
