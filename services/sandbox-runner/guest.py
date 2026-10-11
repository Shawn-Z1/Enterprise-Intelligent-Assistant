"""只在独立Docker任务容器执行的固定文件检查与代码入口；应用后端不导入此模块。"""
import csv
import io
import json
import math
import os
import sys
import zipfile
import subprocess
from pathlib import Path

import pandas as pd
import openpyxl

MAX_CELLS=1_000_000
MAX_ROWS=200_000
MAX_TEXT=2_000_000


def sample_rows(frame, count):
    """接收表格与样例行数，返回最多20列和有限字符串的真实样例；不冒充完整表格。"""
    rows=json.loads(frame.iloc[:count,:20].to_json(orient='records',date_format='iso',force_ascii=False))
    return [{key:value[:256] if isinstance(value,str) else value for key,value in row.items()} for row in rows]


def validate_zip(path, kind):
    """接收DOCX/XLSX路径及类型，限制解压和关系；普通超链接不加载，宏/外部内容拒绝。"""
    from xml.etree import ElementTree
    with zipfile.ZipFile(path) as archive:
        entries=archive.infolist()
        required='word/document.xml' if kind=='docx' else 'xl/workbook.xml'
        if required not in archive.namelist() or len(entries)>2000 or sum(x.file_size for x in entries)>100*1024*1024:
            raise ValueError('压缩文档格式或解压规模异常')
        for entry in entries:
            if entry.flag_bits&1 or entry.file_size>max(1024*1024,entry.compress_size*200) or 'vba' in entry.filename.lower():
                raise ValueError('不支持宏、加密或异常压缩文档')
            if entry.filename.endswith('.rels'):
                if entry.file_size>1024*1024:raise ValueError('关系文件超限')
                for relation in ElementTree.fromstring(archive.read(entry)):
                    if relation.get('TargetMode')=='External' and not relation.get('Type','').endswith('/hyperlink'):
                        raise ValueError('不加载文档外部内容')


def text_units(item):
    """接收固定附件清单项，在客体解析文字单元和结构；返回单元、标题、无文字页和格式限制。"""
    path=Path('/input')/item['path'];kind=item['format'];headings=[];empty_pages=[]
    if kind in ('txt','md'):
        raw=path.read_bytes()
        if b'\x00' in raw:raise ValueError('文本含二进制内容')
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:text=raw.decode('gb18030')
        units=text.splitlines()
        if kind=='md':headings=[line[:160] for line in units if line.lstrip().startswith('#')][:30]
    elif kind=='docx':
        validate_zip(path,kind)
        from docx import Document
        from docx.table import Table
        units=[]
        for block in Document(path).iter_inner_content():
            if isinstance(block,Table):
                units.extend('\t'.join(cell.text for cell in row.cells) for row in block.rows)
            else:
                units.append(block.text)
                if block.style and block.style.name.startswith('Heading'):headings.append(block.text[:160])
    elif kind=='pdf':
        from pypdf import PdfReader
        if not path.read_bytes()[:8].startswith(b'%PDF-'):raise ValueError('PDF真实格式不符')
        reader=PdfReader(path)
        if reader.is_encrypted:raise ValueError('暂不支持加密PDF')
        if len(reader.pages)>300:raise ValueError('PDF最多300页，请缩小范围')
        units=[]
        for index,page in enumerate(reader.pages):
            value=page.extract_text() or '';units.append(value)
            if not value.strip():empty_pages.append(index+1)
            if sum(map(len,units))>MAX_TEXT:raise ValueError('PDF文本规模超过200万字符')
        if not any(value.strip() for value in units):raise ValueError('PDF没有可提取文字；暂不支持扫描识别/OCR')
    else:raise ValueError('该格式不是文字文档')
    text='\n'.join(units)
    if len(text)>MAX_TEXT:raise ValueError('文档超过200万字符，请缩小范围')
    return units,headings[:30],empty_pages


def prepare_files(files):
    """接收本次附件，在受限实例固定解析概况；逐文件失败可见，返回长度/结构/样例与覆盖标志。"""
    summaries=[]
    for index,item in enumerate(files):
        base={'attachment_id':item['attachment_id'],'file':item['name'],'format':item['format']}
        try:
            if item['format'] in ('csv','xlsx'):
                # 保留原输入序号，以便概况表名与后续代码的tables映射一致。
                frames,metadata=load_frames(files,{item['attachment_id']})
                rows=[row for row in metadata if row['attachment_id']==item['attachment_id']]
                for row in rows:
                    row['columns_total']=len(row['columns']);row['columns_truncated']=len(row['columns'])>100
                    row['columns']=row['columns'][:100];row['types']={k:v for k,v in row['types'].items() if k in row['columns']}
                    row['missing']={k:v for k,v in row['missing'].items() if k in row['columns']}
                summaries.extend(rows)
            else:
                units,headings,empty_pages=text_units(item);text='\n'.join(units)
                complete=len(text)<=1800
                summaries.append(base|{'status':'ready','text_length':len(text),'unit_count':len(units),
                    'page_count':len(units) if item['format']=='pdf' else None,'headings':headings,'empty_pages':empty_pages,
                    'sample':text if complete else text[:1400],'sample_complete':complete,
                    'coverage':{'start':0,'end':len(text) if complete else min(1400,len(text)),'total':len(text)}})
        except Exception as error:summaries.append(base|{'status':'failed','error':str(error)[:600],'sample_complete':False})
    return {'files':summaries,'message':'附件预处理完成；概况和样例不是已读全文','artifacts':[]}


def read_file(files, options):
    """接收本人附件与有界读取参数，返回真实正文/表格和续读位置；只在独立任务容器解析。"""
    item=next((x for x in files if x['attachment_id']==options['attachment_id']),None)
    if item is None:raise ValueError('附件不属于本次输入')
    if item['format'] in ('csv','xlsx'):
        frames,metadata=load_frames(files,{item['attachment_id']});table=options.get('table')
        candidates=[row for row in metadata if row['attachment_id']==item['attachment_id']]
        selected=next((x for x in candidates if x['table']==table),None) if table else candidates[0]
        if selected is None:raise ValueError('工作表不属于指定附件')
        frame=frames[selected['table']];start=int(options.get('offset',0));count=min(int(options.get('limit',20)),20)
        column_start=int(options.get('column_offset',0));column_end=min(column_start+int(options.get('column_limit',20)),len(frame.columns))
        if start>len(frame) or column_start>=len(frame.columns):raise ValueError('读取位置超过表格范围')
        # 详细读取不截断单元格；用行/列分页和响应字符预算控制规模。
        records=json.loads(frame.iloc[start:start+count,column_start:column_end].to_json(orient='records',date_format='iso',force_ascii=False))
        while records and len(json.dumps(records,ensure_ascii=False))>12000:
            if len(records)==1:raise ValueError('单行读取超过预算，请减小column_limit')
            records.pop()
        end=min(start+len(records),len(frame));visible=list(frame.columns[column_start:column_end])
        selected={**selected,'columns':visible,'sample':sample_rows(frame.iloc[start:,column_start:column_end],3),
            'types':{k:v for k,v in selected['types'].items() if k in visible},'missing':{k:v for k,v in selected['missing'].items() if k in visible}}
        return {'files':[selected],'rows':records,'columns':list(frame.columns[column_start:column_end]),'row_count':len(frame),
            'coverage':{'scope':selected['table'],'unit':'rows','start':start,'end':end,'total':len(frame),
                'column_start':column_start,'column_end':column_end,'columns_total':len(frame.columns),
                'complete':start==0 and end==len(frame) and column_start==0 and column_end==len(frame.columns)},
            'has_more':end<len(frame) or column_end<len(frame.columns),'has_more_rows':end<len(frame),'has_more_columns':column_end<len(frame.columns),
            'next_offset':end,'next_column_offset':column_end,'truncated':False,'artifacts':[]}
    units,headings,empty_pages=text_units(item);text='\n'.join(units);page=options.get('page')
    if page:
        if item['format']!='pdf' or page>len(units):raise ValueError('页码不适用于该文档')
        text=units[page-1]
    start=int(options.get('offset',0));limit=min(int(options.get('limit',12000)),12000)
    if start>len(text):raise ValueError('读取位置超过文字范围')
    end=min(start+limit,len(text))
    return {'files':[{'attachment_id':item['attachment_id'],'file':item['name'],'format':item['format'],'headings':headings,'empty_pages':empty_pages}],
        'text':text[start:end],'coverage':{'scope':f'page:{page}' if page else 'full_document','unit':'characters',
            'start':start,'end':end,'total':len(text),'complete':not page and start==0 and end==len(text) and not empty_pages},
        'has_more':end<len(text),'next_offset':end,'artifacts':[]}


def load_frames(files, selected=None):
    """接收固定输入清单，在隔离客体验证并解析CSV/XLSX；返回数据框和真实检查元数据。"""
    frames,metadata={},[]
    cells=0
    for index,item in enumerate(files):
        if item['format'] not in ('csv','xlsx') or selected is not None and item['attachment_id'] not in selected:continue
        path=Path('/input')/item['path']
        sheets=[]
        if item['format']=='csv':
            raw=path.read_bytes()
            try:text=raw.decode('utf-8-sig')
            except UnicodeDecodeError:text=raw.decode('gb18030')
            sheets=[('CSV',csv.reader(io.StringIO(text)))]
        else:
            # 普通超链接只保留单元格文本，不访问目标；复用解压/宏及外部加载边界检查。
            validate_zip(path,'xlsx')
            workbook=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
            sheets=[(sheet.title,sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
        for sheet_index,(sheet_name,iterator) in enumerate(sheets):
            iterator=iter(iterator)
            header=next(iterator,[])
            columns=[str(v) if v is not None else f'column_{i+1}' for i,v in enumerate(header)]
            if not columns or len(columns)>2000 or len(set(columns))!=len(columns):raise ValueError('空表、重复列名或列数超限')
            rows=[]
            for row in iterator:
                cells+=len(columns)
                if cells>MAX_CELLS or len(rows)>=MAX_ROWS:raise ValueError('文件超过100万单元格或20万行限制')
                values=list(row)
                if any(isinstance(v,str) and v.startswith('=') for v in values) and item['format']=='xlsx':raise ValueError('不执行Excel公式，请先导出为值')
                if any(isinstance(v,str) and len(v)>10000 for v in values):raise ValueError('单元格文本超限')
                rows.append((values+[None]*len(columns))[:len(columns)])
            frame=pd.DataFrame(rows,columns=columns)
            for column in columns:
                converted=pd.to_numeric(frame[column].replace('',None),errors='coerce')
                if converted.notna().sum()==frame[column].replace('',None).notna().sum():frame[column]=converted
            table=f'f{item.get("index",index+1)}_s{sheet_index+1}'
            frames[table]=frame
            metadata.append({'attachment_id':item['attachment_id'],'file':item['name'],'sheet':sheet_name,'table':table,'columns':columns,
                             'types':{c:str(frame[c].dtype) for c in columns},'row_count':len(frame),'missing':{c:int(frame[c].isna().sum()) for c in columns},
                             'sample':sample_rows(frame,3),'sample_complete':len(frame)<=3 and len(columns)<=20,'status':'ready','format':item['format']})
    return frames,metadata


def emit_result(rows,chart=None,description=''):
    """接收模型代码计算出的记录、图表字段和口径，保存JSON；真实图表及安全CSV由固定入口生成。"""
    if isinstance(rows,pd.DataFrame):rows=json.loads(rows.to_json(orient='records',date_format='iso',force_ascii=False))
    Path('/output/result.json').write_text(json.dumps({'rows':rows,'chart':chart or {},'description':description},ensure_ascii=False,allow_nan=False),encoding='utf-8')


def finish_result(metadata):
    """验证客体结果结构并导出固定CSV/PNG；返回有限模型观察，完整有限产物由节点收集。"""
    path=Path('/output/result.json')
    if not path.exists() or path.stat().st_size>2*1024*1024:raise ValueError('代码须调用emit_result，结果不得超过2MiB')
    value=json.loads(path.read_text(encoding='utf-8'))
    rows=value.get('rows')
    if not isinstance(rows,list) or len(rows)>200000 or any(not isinstance(r,dict) for r in rows):raise ValueError('结果必须为记录列表')
    columns=list(rows[0]) if rows else []
    if len(columns)>100 or any(set(r)!=set(columns) for r in rows):raise ValueError('结果字段不一致或超限')
    for row in rows:
        if any(not isinstance(v,(str,int,float,bool,type(None))) or isinstance(v,float) and not math.isfinite(v) for v in row.values()):raise ValueError('结果单元格须为有限标量')
    with Path('/output/result.csv').open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.writer(stream)
        writer.writerow(["'"+c if c.startswith(('=','+','-','@')) else c for c in columns])
        for row in rows:writer.writerow(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')) else v for v in row.values()])
    chart=value.get('chart') or {}
    points=[]
    if chart.get('type') in ('bar','line'):
        x,y=chart.get('x'),chart.get('y')
        if x not in columns or y not in columns:raise ValueError('图表字段不在结果中')
        for row in rows[:20]:
            if not isinstance(row[y],(float,int)) or isinstance(row[y],bool):raise ValueError('图表纵轴必须为数值')
            points.append({'label':str(row[x])[:120],'value':row[y]})
        import matplotlib.pyplot as plt
        plt.rcParams['font.sans-serif']=['Noto Sans CJK JP','DejaVu Sans']
        plt.rcParams['axes.unicode_minus']=False
        fig,ax=plt.subplots(figsize=(8,4))
        labels=[p['label'] for p in points];numbers=[p['value'] for p in points]
        (ax.bar if chart['type']=='bar' else ax.plot)(labels,numbers)
        ax.set_title(str(chart.get('title',''))[:120])
        ax.set_ylabel(str(chart.get('unit',''))[:40])
        ax.tick_params(axis='x',labelrotation=30)
        fig.tight_layout();fig.savefig('/output/chart.png');plt.close(fig)
    return {'rows':rows[:200],'columns':columns,'row_count':len(rows),'truncated':len(rows)>200,'chart':{**chart,'points':points,'partial':len(rows)>20} if points else None,
            'description':str(value.get('description',''))[:2000],'files':metadata,'artifacts':['result.csv']+(['chart.png'] if points else [])}


def emit_text(text, description=''):
    """接收实际计算文字和说明，保存有限正文；完整文件由代码写入OUTPUT_DIR。"""
    Path('/output/text_result.json').write_text(json.dumps({'text':str(text)[:12000],'description':str(description)[:2000]},ensure_ascii=False),encoding='utf-8')


def convert_document(source, target='pdf', name=None):
    """接收本任务文档路径、DOCX/PDF目标及可选名称，调用预装转换器并返回实际文件路径。"""
    import re
    source=Path(source).resolve()
    if not source.is_file() or not any(source.is_relative_to(root) for root in (Path('/input'),Path('/output'))):raise ValueError('转换输入必须属于本任务')
    if target not in ('docx','pdf'):raise ValueError('转换目标仅DOCX/PDF')
    stem=name or source.stem
    if not re.fullmatch(r'[^/\\\x00-\x1f]{1,100}',stem) or stem in ('.','..'):raise ValueError('输出文件名非法')
    docx=source
    if source.suffix.lower() in ('.md','.txt'):
        docx=Path('/output')/(stem+'.docx')
        subprocess.run(['pandoc',str(source),'-f','markdown','--reference-doc=/opt/reference.docx','-o',str(docx)],check=True,timeout=30)
    elif source.suffix.lower()!='.docx':raise ValueError('转换输入仅MD/TXT/DOCX')
    validate_zip(docx,'docx')
    if target=='docx':
        destination=Path('/output')/(stem+'.docx')
        if docx!=destination:
            import shutil
            shutil.copyfile(docx,destination)
        return destination
    subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo-profile','--headless','--norestore','--convert-to','pdf','--outdir','/output',str(docx)],check=True,timeout=40)
    actual=Path('/output')/(docx.stem+'.pdf');destination=Path('/output')/(stem+'.pdf')
    if not actual.is_file():raise ValueError('转换引擎没有生成PDF')
    if actual!=destination:actual.rename(destination)
    return destination


def validate_output(path):
    """接收本任务产物路径，在受限客体核验可打开格式/规模；不在控制器解析复杂文档。"""
    kind=path.suffix.lower()
    if kind=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted or not 0<len(reader.pages)<=300:raise ValueError('PDF产物无有效页面或超限')
    elif kind in ('.docx','.xlsx'):
        validate_zip(path,kind[1:])
        if kind=='.docx':
            from docx import Document
            Document(path)
        else:
            workbook=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
            cells=0
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    cells+=len(row)
                    if cells>MAX_CELLS:raise ValueError('XLSX产物超过单元格上限')
                    if any(cell.data_type=='f' for cell in row):raise ValueError('XLSX产物不能包含未计算公式')
            workbook.close()
    elif kind in ('.txt','.md'):path.read_text(encoding='utf-8-sig')


def main():
    """读取固定输入清单，执行预处理/读取/自动代码并保存回执；每任务使用独立客体。"""
    request=json.loads(Path('/input/request.json').read_text(encoding='utf-8'))
    try:
        operation=request['operation'];files=request['files']
        if operation in ('prepare','inspect'):result=prepare_files(files)
        elif operation=='read':result=read_file(files,request['read'])
        else:
            frames,metadata=load_frames(files)
            if request['language']=='python':
                import numpy as np
                import matplotlib.pyplot as plt
                plt.rcParams['font.sans-serif']=['Noto Sans CJK JP','DejaVu Sans']
                plt.rcParams['axes.unicode_minus']=False
                namespace={'pd':pd,'np':np,'tables':frames,'files':{item['attachment_id']:Path('/input')/item['path'] for item in files},
                    'attachments':files,'OUTPUT_DIR':Path('/output'),'emit_result':emit_result,'emit_text':emit_text,
                    'convert_document':convert_document,'__name__':'__main__'}
                # 任意生成代码只在无网络、无凭证客体内执行，禁止应用侧导入此执行器。
                exec(compile(request['code'],'<sandbox-code>','exec'),namespace)
            else:
                import duckdb
                import sqlglot
                from sqlglot import exp
                statements=sqlglot.parse(request['code'],read='duckdb')
                if len(statements)!=1 or not isinstance(statements[0],(exp.Select,exp.Union)):raise ValueError('文件SQL只允许单条只读查询')
                connection=duckdb.connect(config={'enable_external_access':False,'autoload_known_extensions':False,'autoinstall_known_extensions':False,'threads':2,'memory_limit':'512MB'})
                for name,frame in frames.items():connection.register(name,frame)
                emit_result(connection.execute(request['code']).fetchdf(),request.get('chart',{}),'DuckDB只查询本任务文件表')
            result=finish_result(metadata) if Path('/output/result.json').is_file() else {'files':[],'description':'代码已实际执行，输出以回执/产物为准'}
            if Path('/output/text_result.json').is_file():result.update(json.loads(Path('/output/text_result.json').read_text(encoding='utf-8')))
            internal={'.ready','receipt.json','result.json','text_result.json'}
            artifacts=[]
            for path in Path('/output').iterdir():
                if path.name in internal:continue
                if path.is_symlink() or not path.is_file():raise ValueError('输出只允许顶层普通文件')
                if path.suffix.lower() not in ('.txt','.md','.csv','.xlsx','.docx','.pdf','.png'):raise ValueError('产物格式暂不支持')
                validate_output(path)
                artifacts.append(path.name)
            result['artifacts']=artifacts
        receipt={'ok':True,'result':result}
    except Exception as error:receipt={'ok':False,'error':{'code':'SANDBOX_CODE_ERROR','message':str(error)[:2000]}}
    Path('/output/receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,allow_nan=False),encoding='utf-8')
    Path('/output/.ready').write_text('ready')
    import time
    while True:time.sleep(1)


if __name__=='__main__':main()
